"""The Stratification tab's exports: the chart as a vector PDF, its data as an Excel workbook."""

from __future__ import annotations

import io
import math

import numpy as np
import pytest
from fastapi.testclient import TestClient

from sbfsem.analysis.boundary import BoundarySurface, IPLBoundaries
from sbfsem.analysis.stratification import annotation_depths, ipl_depth
from sbfsem.export.stratification import (
    StrataCell,
    StrataSettings,
    annotation_rows_per_page,
    ipl_ticks,
    stratification_pdf,
    stratification_xlsx,
)
from sbfsem.io.legacy_json import read_legacy_json
from sbfsem_web.app import create_app


@pytest.fixture(scope="module")
def flat():
    """A flat IPL between z = 130 and z = 150, covering c943's whole extent."""
    pts = np.linspace(0, 400, 20)
    return IPLBoundaries(
        BoundarySurface("INL", "NeitzInferiorMonkey", np.zeros((3, 3)), pts, pts,
                        np.full((20, 20), 130.0)),
        BoundarySurface("GCL", "NeitzInferiorMonkey", np.zeros((3, 3)), pts, pts,
                        np.full((20, 20), 150.0)),
    )


def _cell(neuron, flat, *, name="c943", color="#d62728", include_soma=True):
    result = ipl_depth(neuron, boundaries=flat, num_bins=20, include_soma=include_soma)
    table = annotation_depths(neuron, boundaries=flat, include_soma=include_soma)
    return StrataCell(name=name, label=f"{name} test", volume="Inferior Monkey", color=color,
                      result=result, annotations=table)


def test_the_gridlines_run_from_the_inl_to_the_gcl_every_step():
    assert ipl_ticks(20) == [(0.0, "INL"), (0.2, "20"), (0.4, "40"), (0.6, "60"), (0.8, "80"), (1.0, "GCL")]
    assert [label for _, label in ipl_ticks(25)] == ["INL", "25", "50", "75", "GCL"]
    # A step that does not divide 100 ends with a shorter interval; the GCL is always a line.
    assert [label for _, label in ipl_ticks(30)] == ["INL", "30", "60", "90", "GCL"]
    assert [label for _, label in ipl_ticks(12.5)][:3] == ["INL", "12.5", "25"]
    assert [label for _, label in ipl_ticks(100)] == ["INL", "GCL"]
    assert ipl_ticks(float("nan")) == ipl_ticks(20)


def test_labels_are_thinned_as_the_page_thins_them():
    """The PDF and the page label the same gridlines (charts.js labeledTicks)."""
    from sbfsem.export.stratification import labeled_ticks

    assert labeled_ticks(ipl_ticks(20), 86, 26) == [0, 1, 2, 3, 4, 5]
    assert [ipl_ticks(30)[i][1] for i in labeled_ticks(ipl_ticks(30), 43, 26)] == ["INL", "30", "60", "GCL"]
    fine = labeled_ticks(ipl_ticks(2), 8.6, 26)
    assert fine[0] == 0 and fine[-1] == 50
    assert all((b - a) * 8.6 >= 26 - 1e-9 for a, b in zip(fine, fine[1:], strict=False))


def test_the_workbook_has_a_summary_column_per_cell_and_a_sheet_per_cell(c943, flat):
    from openpyxl import load_workbook

    a = _cell(c943, flat)
    b = _cell(c943, flat, name="c944", color="#1f77b4", include_soma=False)
    book = load_workbook(io.BytesIO(stratification_xlsx([a, b], StrataSettings(version="9.9.9"))))
    assert book.sheetnames == ["Summary", "c943 Inferior Monkey", "c944 Inferior Monkey"]
    summary = [[cell.value for cell in row] for row in book["Summary"].iter_rows()]
    head = next(r for r in summary if r and r[0] == "bin")
    assert head[:5] == ["bin", "c943 test center (%)", "c943 test count", "c944 test center (%)", "c944 test count"]
    start = summary.index(head) + 1
    rows = summary[start:start + 20]
    assert [r[0] for r in rows] == list(range(1, 21))
    assert [r[2] for r in rows] == [int(v) for v in a.result.counts]
    assert [r[4] for r in rows] == [int(v) for v in b.result.counts]
    assert rows[0][1] == pytest.approx(a.result.centers[0] * 100, abs=0.01)
    stats = {r[0]: r for r in summary if r and r[0] in ("annotations counted", "median (%)", "soma left out")}
    assert stats["annotations counted"][1] == a.result.stats.n
    assert stats["median (%)"][1] == pytest.approx(a.result.stats.median * 100, abs=0.01)
    assert stats["soma left out"][3] > 0 and stats["soma left out"][1] == 0
    assert any(r and r[0] == "SBFSEM-tools" and r[1] == "9.9.9" for r in summary)

    sheet = [[cell.value for cell in row] for row in book["c944 Inferior Monkey"].iter_rows()]
    assert sheet[0][0] == "c944 test · Inferior Monkey · depth along Z (MATLAB's column)"
    assert sheet[1] == ["Location ID", "section", "X (µm)", "Y (µm)", "Z (µm)", "radius (µm)",
                        "IPL depth (%)", "depth (µm)", "IPL thickness (µm)", "counted", "left out because"]
    body = sheet[2:]
    assert len(body) == len(c943.get_cell_nodes())
    assert sum(1 for r in body if r[9]) == b.result.stats.n
    assert {r[10] for r in body if not r[9]} == {"soma"}
    # The flat IPL is 20 µm thick everywhere: the depth in µm is the fraction of it.
    inside = [r for r in body if r[6] is not None]
    assert inside and all(r[8] == pytest.approx(20.0) for r in inside)
    assert all(r[7] == pytest.approx(r[6] / 100 * 20.0, abs=1e-3) for r in inside)
    assert any(r and r[0] == "depth measured" and r[1].startswith("depth along Z") for r in summary)


def test_the_pdf_is_the_chart_as_vectors_then_the_data_when_asked(c943, flat):
    from pypdf import PdfReader

    cell = _cell(c943, flat)
    settings = StrataSettings(tick_pct=20, version="9.9.9")
    chart = PdfReader(io.BytesIO(stratification_pdf([cell], settings)))
    assert len(chart.pages) == 1
    page = chart.pages[0]
    assert not page.images, "the chart is drawn as vectors, not a picture"
    text = page.extract_text()
    for word in ("INL", "20", "40", "60", "80", "GCL", "IPL depth (%)", "c943 test", "gridlines every 20 %"):
        assert word in text, word
    width, height = float(page.mediabox.width), float(page.mediabox.height)
    assert (width, height) == (8.0 * 72, 4.6 * 72), "the page is the size of the figure"
    assert chart.metadata.title == "IPL stratification"
    assert "Claude" not in str(chart.metadata) and "Anthropic" not in str(chart.metadata)

    stats = PdfReader(io.BytesIO(stratification_pdf([cell], settings, statistics=True)))
    assert len(stats.pages) == 2
    assert "Binned counts" in stats.pages[1].extract_text()

    every = PdfReader(io.BytesIO(stratification_pdf([cell], settings, statistics=True, annotations=True)))
    n = len(c943.get_cell_nodes())
    assert len(every.pages) == 2 + math.ceil(n / annotation_rows_per_page())
    first_id = str(int(c943.get_cell_nodes()["ID"].iloc[0]))
    assert first_id in every.pages[2].extract_text()


@pytest.fixture
def client(i943_path, flat):
    app = create_app()
    with TestClient(app) as test_client:
        svc = app.state.service
        svc.offline = True
        svc.add_neuron(read_legacy_json(i943_path))
        svc.boundaries = lambda source: flat
        yield test_client


def test_the_routes_send_the_files_for_the_copies_named(client):
    body = {"cells": [{"source": "i", "id": 943, "label": "c943 ON", "color": "#2ca02c"}],
            "num_bins": 12, "tick_pct": 25, "statistics": False}
    pdf = client.post("/api/stratification/export.pdf", json=body)
    assert pdf.status_code == 200, pdf.text
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")
    assert 'filename="stratification-c943.pdf"' in pdf.headers["content-disposition"]

    xlsx = client.post("/api/stratification/export.xlsx", json=body)
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.content[:2] == b"PK"
    from openpyxl import load_workbook

    summary = [[c.value for c in r] for r in load_workbook(io.BytesIO(xlsx.content))["Summary"].iter_rows()]
    head = next(r for r in summary if r and r[0] == "bin")
    assert head[1] == "c943 ON center (%)"
    assert sum(1 for r in summary[summary.index(head) + 1:] if isinstance(r[0], int)) == 12

    # Every annotation's depth past the PDF's limit is refused, saying where they all are.
    import sbfsem.export.stratification as export

    limit = export.PDF_MAX_ANNOTATION_ROWS
    try:
        export.PDF_MAX_ANNOTATION_ROWS = 10
        over = client.post("/api/stratification/export.pdf", json={**body, "annotations": True})
        assert over.status_code == 400 and "Data .xlsx" in over.json()["detail"]
    finally:
        export.PDF_MAX_ANNOTATION_ROWS = limit

    # A cell not loaded is a 404; a bad color, a tick step under 1 % and no cells are refused.
    missing = client.post("/api/stratification/export.xlsx",
                          json={**body, "cells": [{"source": "i", "id": 12345}]})
    assert missing.status_code == 404
    for bad in ({**body, "cells": [{**body["cells"][0], "color": "red"}]}, {**body, "tick_pct": 0.5},
                {**body, "cells": []}):
        assert client.post("/api/stratification/export.pdf", json=bad).status_code == 422
