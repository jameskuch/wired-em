"""IPL depth measured along the layers' normals, beside MATLAB's column along Z.

James Anderson's request of 2026-10-22: depth by projecting onto each fitted
layer, and the IPL's thickness in µm.  The library and the bare API keep the
column (ground rule 1); the viewer measures along the normals.  The analytic
cases here know both answers exactly.
"""

from __future__ import annotations

import inspect

import numpy as np
import pytest
from fastapi.testclient import TestClient

from sbfsem.analysis import stratification
from sbfsem.analysis.boundary import (
    IPL_METHODS,
    BoundarySurface,
    IPLBoundaries,
    load_ipl_boundaries,
    thickness_summary,
)
from sbfsem.geometry import meshops
from sbfsem.service import render_schema
from sbfsem.service.facade import SbfsemService

SOURCE = "NeitzInferiorMonkey"


def _surface(layer: str, height, *, lo: float = 0.0, hi: float = 200.0, n: int = 201) -> BoundarySurface:
    """A boundary whose fit is ``height(x, y)`` sampled on an n x n grid (rows are Y)."""
    pts = np.linspace(lo, hi, n)
    x, y = np.meshgrid(pts, pts)
    return BoundarySurface(layer, SOURCE, np.zeros((3, 3)), pts, pts, height(x, y).astype(float))


def _points(boundaries: IPLBoundaries, fractions, xy) -> np.ndarray:
    """Points between the two fits by the column's fraction, at each XY."""
    xy = np.asarray(xy, dtype=float)
    z_inl = boundaries.inl.xy_eval(xy[:, 0], xy[:, 1])
    z_gcl = boundaries.gcl.xy_eval(xy[:, 0], xy[:, 1])
    out = [np.column_stack([xy, z_inl + f * (z_gcl - z_inl)]) for f in fractions]
    return np.concatenate(out)


XY = np.array([[60.0, 70.0], [100.0, 100.0], [140.0, 90.0], [80.0, 130.0]])
FRACTIONS = (-0.1, 0.0, 0.25, 0.5, 0.9, 1.0, 1.15)


# -- tilted planes ------------------------------------------------------------------
def test_parallel_tilted_planes_agree_on_the_fraction_and_differ_on_the_thickness_by_cos():
    """Two parallel planes tilted 20°: the fraction is the same along Z and along
    the normal, at every depth; the thickness along Z is T, along the normal
    T cos θ, and the depth in µm scales the same way."""
    slope = np.tan(np.radians(20.0))
    gap = 30.0
    b = IPLBoundaries(_surface("INL", lambda x, y: 10 + slope * x),
                      _surface("GCL", lambda x, y: 10 + gap + slope * x))
    pts = _points(b, FRACTIONS, XY)
    column = b.depth(pts, "column")
    normal = b.depth(pts, "normal")
    np.testing.assert_allclose(normal.fraction, column.fraction, atol=1e-9)
    np.testing.assert_allclose(column.thickness_um, gap, rtol=1e-12)
    np.testing.assert_allclose(normal.thickness_um, gap * np.cos(np.radians(20.0)), rtol=1e-9)
    np.testing.assert_allclose(normal.depth_um, normal.fraction * normal.thickness_um, atol=1e-9)
    np.testing.assert_allclose(column.depth_um, column.fraction * gap, atol=1e-9)


def test_planes_that_are_not_parallel_give_the_exact_distances_to_each():
    """A flat INL at z = 0 and a GCL tilted to z = 20 + 0.2 x: the depth along the
    normal is the point's height above the INL, the distance to the GCL is to the
    tilted plane, and the column's fraction differs from it."""
    b = IPLBoundaries(_surface("INL", lambda x, y: 0 * x),
                      _surface("GCL", lambda x, y: 20 + 0.2 * x))
    pts = _points(b, (0.1, 0.5, 0.9), XY)
    normal = b.depth(pts, "normal")
    x, z = pts[:, 0], pts[:, 2]
    d_inl = z
    d_gcl = (20 + 0.2 * x - z) / np.sqrt(1 + 0.2 ** 2)
    np.testing.assert_allclose(normal.depth_um, d_inl, atol=1e-9)
    np.testing.assert_allclose(normal.thickness_um, d_inl + d_gcl, atol=1e-9)
    np.testing.assert_allclose(normal.fraction, d_inl / (d_inl + d_gcl), atol=1e-12)
    column = b.micron2ipl(pts)
    np.testing.assert_allclose(column, z / (20 + 0.2 * x), atol=1e-12)
    assert np.abs(column - normal.fraction).max() > 0.004  # half a percent of the IPL at 11°


def test_a_gcl_below_the_inl_still_counts_depth_from_the_inl():
    """The Neitz volumes have the GCL at smaller Z than the INL."""
    b = IPLBoundaries(_surface("INL", lambda x, y: 80 + 0.1 * y),
                      _surface("GCL", lambda x, y: 50 + 0.1 * y))
    assert b.gcl_side == -1.0
    pts = _points(b, (0.0, 0.3, 1.0), XY)
    normal = b.depth(pts, "normal")
    np.testing.assert_allclose(normal.fraction, np.repeat([0.0, 0.3, 1.0], len(XY)), atol=1e-9)
    assert (normal.thickness_um > 0).all()


# -- a curved IPL ----------------------------------------------------------------------
def test_concentric_spheres_are_one_thickness_everywhere_along_the_normal():
    """An INL and a GCL on concentric spheres 30 µm apart (lower caps, radius 400
    and 370 µm, centered over the middle): along the normal the IPL is 30 µm
    thick wherever it is measured and a point's fraction is its distance in from
    the INL sphere over 30; along Z it reads thicker away from the axis."""
    center = np.array([100.0, 100.0, 450.0])
    r_inl, r_gcl = 400.0, 370.0

    def cap(radius):
        return lambda x, y: center[2] - np.sqrt(radius ** 2 - (x - center[0]) ** 2 - (y - center[1]) ** 2)

    b = IPLBoundaries(_surface("INL", cap(r_inl)), _surface("GCL", cap(r_gcl)))
    xy = np.array([[100.0, 100.0], [40.0, 60.0], [160.0, 150.0], [30.0, 170.0]])
    pts = _points(b, (0.1, 0.5, 0.9), xy)
    normal = b.depth(pts, "normal")
    rho = np.linalg.norm(pts - center, axis=1)
    np.testing.assert_allclose(normal.thickness_um, r_inl - r_gcl, atol=2e-3)
    np.testing.assert_allclose(normal.depth_um, r_inl - rho, atol=2e-3)
    np.testing.assert_allclose(normal.fraction, (r_inl - rho) / (r_inl - r_gcl), atol=1e-4)
    column = b.depth(pts, "column")
    off_axis = np.hypot(pts[:, 0] - center[0], pts[:, 1] - center[1]) > 50
    assert (column.thickness_um[off_axis] > normal.thickness_um[off_axis] + 0.5).all()


def test_the_nearest_point_on_a_curved_fit_is_the_nearest_of_a_dense_sampling():
    """On a real fit (Inferior Monkey's INL), the solved distance is never more
    than a dense sampling's nearest point, and within the sampling's own error."""
    from scipy.spatial import cKDTree

    surface = load_ipl_boundaries(SOURCE).inl
    rng = np.random.default_rng(3)
    x = rng.uniform(60, 200, 400)
    y = rng.uniform(60, 200, 400)
    pts = np.column_stack([x, y, surface.xy_eval(x, y) - rng.uniform(1, 30, 400)])
    _, distance = surface.nearest(pts)
    assert np.isfinite(distance).all()
    xs = np.linspace(40, 220, 1200)
    gx, gy = np.meshgrid(xs, xs)
    cloud = np.column_stack([gx.ravel(), gy.ravel(), surface.xy_eval(gx.ravel(), gy.ravel())])
    sampled, _ = cKDTree(cloud).query(pts)
    # The sampling's spacing is 0.15 µm: its nearest point is at most that far off.
    assert (sampled - np.abs(distance) > -0.01).all()
    assert np.median(sampled - np.abs(distance)) < 1e-3


# -- what has no depth ---------------------------------------------------------------
def test_a_point_whose_nearest_layer_point_is_beyond_the_fit_has_no_normal_depth():
    """Steep planes near the fit's edge: straight down the point is inside the
    rectangle, but the nearest point of the layer is past its edge."""
    b = IPLBoundaries(_surface("INL", lambda x, y: 0.8 * x),
                      _surface("GCL", lambda x, y: 30 + 0.8 * x))
    pts = np.array([[2.0, 100.0, 1.6 + 20.0], [100.0, 100.0, 95.0], [np.nan, 1.0, 1.0]])
    normal = b.depth(pts, "normal")
    assert np.isnan(normal.fraction[0]) and np.isnan(normal.thickness_um[0])
    assert np.isfinite(normal.fraction[1])
    assert np.isnan(normal.fraction[2])
    assert np.isfinite(b.micron2ipl(pts[:1]))[0], "the column still reads it"


def test_an_unknown_method_is_refused():
    b = load_ipl_boundaries(SOURCE)
    with pytest.raises(ValueError, match="IPL depth method"):
        b.depth(np.zeros((1, 3)), "vertical")
    with pytest.raises(ValueError, match="IPL depth method"):
        b.micron2ipl(np.zeros((1, 3)), "Normal")


# -- the library's default is MATLAB's ---------------------------------------------
def test_the_library_and_the_api_default_to_the_column(c943):
    """Ground rule 1: every entry point that measures IPL depth defaults to the
    column, and does what the column does when not told otherwise.  Making the
    normal the library's default would change every published number made with
    the library: it fails here."""
    entry_points = [
        IPLBoundaries.micron2ipl, IPLBoundaries.depth, IPLBoundaries.strata_cdata,
        stratification.ipl_depth, stratification.annotation_depths,
        stratification.synapse_stratification, stratification.branch_stratification,
        meshops.clip_mesh_by_stratification, meshops.flatten_render,
    ]
    for fn in entry_points:
        assert inspect.signature(fn).parameters["method"].default == "column", fn.__qualname__
    for fn in (SbfsemService.render_mesh, SbfsemService.render_buffers, SbfsemService.clipped_mesh,
               SbfsemService.stratification, SbfsemService.stratification_tables,
               SbfsemService.synapse_stratification):
        assert inspect.signature(fn).parameters["ipl_method"].default == "column", fn.__qualname__
    assert render_schema.field_by_key("ipl_method").default == "column"
    assert render_schema.field_by_key("ipl_method").app_value() == "normal"
    assert IPL_METHODS == ("column", "normal")

    # And by behavior: non-parallel planes, where the two differ.
    b = IPLBoundaries(_surface("INL", lambda x, y: 0 * x, hi=400),
                      _surface("GCL", lambda x, y: 20 + 0.2 * x, hi=400))
    pts = _points(b, (0.2, 0.7), XY)
    np.testing.assert_array_equal(b.micron2ipl(pts), b.depth(pts, "column").fraction)
    assert not np.allclose(b.micron2ipl(pts), b.micron2ipl(pts, "normal"))
    result = stratification.ipl_depth(c943, boundaries=b, omit_outliers=False)
    assert result.meta["method"] == "column"
    xyz = c943.get_cell_nodes()[["Xum", "Yum", "Zum"]].to_numpy(dtype=float)
    expected = b.micron2ipl(xyz, "column")
    np.testing.assert_array_equal(np.sort(result.depth), np.sort(expected[np.isfinite(expected)]))


def test_the_annotation_table_carries_depth_and_thickness_in_um(c943):
    b = IPLBoundaries(_surface("INL", lambda x, y: 100 + 0.2 * x, hi=400),
                      _surface("GCL", lambda x, y: 160 + 0.2 * x, hi=400))
    for method in IPL_METHODS:
        table = stratification.annotation_depths(c943, boundaries=b, method=method)
        assert {"depth_um", "ipl_thickness_um"} <= set(table.columns)
        inside = table[np.isfinite(table["ipl_depth"])]
        assert len(inside)
        np.testing.assert_allclose(inside["depth_um"], inside["ipl_depth"] * inside["ipl_thickness_um"],
                                   atol=1e-9)
        want = 60.0 if method == "column" else 60.0 / np.sqrt(1 + 0.2 ** 2)
        np.testing.assert_allclose(inside["ipl_thickness_um"], want, rtol=1e-9)


# -- the volume's thickness ---------------------------------------------------------
def test_the_thickness_summary_of_tilted_parallel_planes():
    slope = np.tan(np.radians(15.0))
    b = IPLBoundaries(_surface("INL", lambda x, y: 10 + slope * y),
                      _surface("GCL", lambda x, y: 50 + slope * y))
    # Markers everywhere, so every column is near both layers.
    grid = np.array([[x, y, 0.0] for x in range(0, 201, 20) for y in range(0, 201, 20)])
    for s in (b.inl, b.gcl):
        s.marker_locations = grid / 0.0075  # Viking pixels at 7.5 nm
    summary = thickness_summary(b, grid=20)
    assert summary["near_columns"] > 0
    assert summary["normal_um"]["median"] == pytest.approx(40 * np.cos(np.radians(15.0)), rel=1e-6)
    assert summary["column_um"]["median"] == pytest.approx(40.0, rel=1e-9)
    assert summary["column_over_normal"]["median"] == pytest.approx(1 / np.cos(np.radians(15.0)), rel=1e-6)
    assert summary["inl_tilt_deg"]["median"] == pytest.approx(15.0, abs=1e-6)
    assert summary["normals_angle_deg"]["max"] == pytest.approx(0.0, abs=1e-5)
    assert len(summary["map"]["thickness_um"]) == 20


def test_each_marked_volume_has_a_thickness_and_its_column_reads_thicker():
    """The five volumes with markers (2026-10-22 numbers: the column 0.1 to 2.4 %
    thicker at the median, up to 15 % at a column)."""
    for source in ("RC1", "RC2", "NeitzInferiorMonkey", "NeitzTemporalMonkey", "NeitzNasalMonkey"):
        summary = thickness_summary(load_ipl_boundaries(source))
        assert summary["near_columns"] > 500, source
        assert 10 < summary["normal_um"]["median"] < 50, source
        assert 1.0 <= summary["column_over_normal"]["median"] < 1.03, source
        assert summary["column_over_normal"]["min"] >= 1.0 - 1e-9, source


# -- the API ----------------------------------------------------------------------
@pytest.fixture(scope="module")
def client(request):
    from sbfsem.io.legacy_json import read_legacy_json
    from sbfsem_web.app import create_app

    fixture = request.path.parents[1] / "fixtures" / "i943.json"
    app = create_app()
    with TestClient(app) as c:
        app.state.service.add_neuron(read_legacy_json(fixture))
        yield c


def test_the_api_measures_along_z_unless_asked_for_the_normals(client):
    column = client.get("/api/neurons/i/943/stratification?omit_outliers=false").json()
    assert column["method"] == "column"
    normal = client.get("/api/neurons/i/943/stratification?omit_outliers=false&ipl_method=normal").json()
    assert normal["method"] == "normal"
    assert normal["stats"]["n"] > 0
    assert client.get("/api/neurons/i/943/stratification?ipl_method=sideways").status_code == 422


def test_the_volume_thickness_route(client):
    reply = client.get("/api/volumes/i/ipl-thickness")
    assert reply.status_code == 200
    body = reply.json()
    assert body["source"] == "NeitzInferiorMonkey"
    assert 20 < body["normal_um"]["median"] < body["column_um"]["median"] < 35
    assert set(body["thinnest"]) == {"x_um", "y_um", "thickness_um", "column_um"}
    assert client.get("/api/volumes/cped/ipl-thickness").status_code == 400
    assert client.get("/api/volumes/i/ipl-thickness?grid=5").status_code == 422


def test_the_strata_exports_say_how_depth_was_measured(client):
    import io

    from openpyxl import load_workbook

    body = {"cells": [{"source": "i", "id": 943}], "omit_outliers": False}
    for method, words in (("column", "depth along Z"), ("normal", "depth along the layers' normals")):
        reply = client.post("/api/stratification/export.xlsx", json={**body, "ipl_method": method})
        assert reply.status_code == 200
        book = load_workbook(io.BytesIO(reply.content))
        line = book["Summary"]["A2"].value
        assert words in line, (method, line)
        sheet = book[book.sheetnames[1]]
        assert words in sheet["A1"].value
    assert client.post("/api/stratification/export.xlsx", json={**body, "ipl_method": "x"}).status_code == 422


def test_the_mesh_route_colors_by_the_method_asked(client):
    base = "/api/neurons/i/943/mesh?color_by_strata=true"
    column = client.get(base).json()
    normal = client.get(base + "&ipl_method=normal").json()
    assert client.get(base + "&ipl_method=bogus").status_code == 422
    a = np.asarray(column["vertex_data"], dtype=float)
    b = np.asarray(normal["vertex_data"], dtype=float)
    assert a.shape == b.shape and a.size
    assert not np.array_equal(a, b)
    assert np.nanmax(np.abs(a - b)) < 0.2
