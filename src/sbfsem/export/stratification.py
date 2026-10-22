"""The Stratification tab's chart and data, as files: a PDF and an Excel workbook.

The PDF's first page is the chart itself, drawn as vectors (matplotlib's PDF
backend, text kept as text), on a page the size of the figure so it can be
placed in a figure as it is, with one line under it saying how it was made.
The pages after it are optional: the statistics and the binned counts, and
each annotation's depth.  The chart is the web page's (``static/js/charts.js``
``stratificationChart``): each cell's counts per bin as a line with a dot per
bin, in the cell's color, a triangle at its median, the margins beyond the IPL
shaded, and gridlines with labels from the INL (0 %) to the GCL (100 %) every
``tick_pct`` (:func:`ipl_ticks`, the same rule as ``charts.js`` ``iplTicks``).

The workbook has a Summary sheet (a column per cell: each bin's center and
count, then the statistics, then the settings) and a sheet per cell with every
annotation's depth and whether the chart counts it
(:func:`sbfsem.analysis.stratification.annotation_depths`).

Both say how the depths were measured (``StrataSettings.method``): along the
layers' normals, or along Z as MATLAB did, so a figure made one way is never
read as the other; each annotation's depth in µm and the IPL's thickness through
it are measured the same way.

Depths are written in percent of the IPL (0 at the INL border, 100 at the GCL
border), as the chart's axis is.  Each cell's bins span its own depths
(``numpy.histogram`` over the cell's range), so bins are matched by number,
not by depth.
"""

from __future__ import annotations

import io
import math
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

import numpy as np
import pandas as pd

from ..analysis.stratification import StratificationResult

__all__ = ["PDF_MAX_ANNOTATION_ROWS", "StrataCell", "StrataSettings", "ipl_ticks", "labeled_ticks", "settings_line", "stratification_pdf",
           "stratification_xlsx", "annotation_rows_per_page"]

#: The chart's x range always shows at least this much beyond the IPL on each side (fractions).
MARGIN_SHOWN = 0.25
#: The most annotation rows the PDF writes (about 360 pages): 1,567 rows took 5.3 s on this
#: Mac (28 pages), in the web service's one process; the workbook has any number.
PDF_MAX_ANNOTATION_ROWS = 20_000


@dataclass
class StrataCell:
    """One cell of the chart: its name (``c162392``), the label the page gave it,
    its volume's name, its color, its histogram and its annotations' depths."""

    name: str
    label: str
    volume: str
    color: str
    result: StratificationResult
    annotations: pd.DataFrame


@dataclass
class StrataSettings:
    """How the chart was made."""

    num_bins: int = 20
    include_soma: bool = True
    omit_outliers: bool = True
    ipl_margin: float = 0.2
    tick_pct: float = 20.0
    #: How the depths were measured (``sbfsem.analysis.boundary.IPL_METHODS``).
    method: str = "column"
    version: str = ""
    made: datetime | None = None


#: How each depth method is named on the chart's line and in the workbook.
METHOD_WORDS = {
    "normal": "depth along the layers' normals",
    "column": "depth along Z (MATLAB's column)",
}


def ipl_ticks(step_pct: float) -> list[tuple[float, str]]:
    """The x axis' gridlines as ``(fraction, label)``: the INL at 0, every
    ``step_pct`` percent up to the GCL at 100, which is always one (a step that
    does not divide 100 ends with a shorter interval).  The same rule as
    ``charts.js`` ``iplTicks``."""
    step = float(step_pct)
    if not math.isfinite(step) or step <= 0:
        step = 20.0
    ticks = [(0.0, "INL")]
    k = 1
    while k * step < 100 - 1e-9:
        value = round(k * step, 6)
        ticks.append((value / 100.0, f"{value:g}"))
        k += 1
    ticks.append((1.0, "GCL"))
    return ticks


def labeled_ticks(ticks: Sequence[tuple[float, str]], gap: float, min_gap: float) -> list[int]:
    """Which gridlines get a label when they are ``gap`` apart: every one when
    there is room, else every k-th from the INL, never one crowding the GCL's.
    The same rule as ``charts.js`` ``labeledTicks``."""
    every = 1 if gap >= min_gap else math.ceil(min_gap / max(gap, 1e-6))
    last = len(ticks) - 1
    step = ticks[1][0] - ticks[0][0]
    out = []
    for i in range(last + 1):
        if i in (0, last) or (i % every == 0 and (ticks[last][0] - ticks[i][0]) * gap / step >= min_gap):
            out.append(i)
    return out


def settings_line(settings: StrataSettings) -> str:
    """One line under the chart: the bins, the soma, the outliers, the ticks, how
    depth was measured, the version and when."""
    parts = [
        f"{settings.num_bins} bins per cell",
        "soma included" if settings.include_soma else "soma left out (annotations over 80 % of the largest radius)",
        f"outliers beyond {settings.ipl_margin * 100:g} % of the IPL omitted" if settings.omit_outliers
        else "outliers kept",
        f"gridlines every {settings.tick_pct:g} %",
        METHOD_WORDS.get(settings.method, settings.method),
    ]
    if settings.version:
        parts.append(f"SBFSEM-tools {settings.version}")
    made = settings.made or datetime.now(UTC)
    parts.append(made.astimezone(UTC).strftime("%Y-%m-%d %H:%M UTC"))
    return " · ".join(parts)


def _pct(value: float | None) -> float | None:
    """A depth as percent, rounded to 0.01 %; None for none."""
    if value is None:
        return None
    value = float(value)
    return round(value * 100.0, 2) if math.isfinite(value) else None


def _stats_rows(cell: StrataCell) -> list[tuple[str, float | int | None]]:
    r = cell.result
    left = cell.annotations["left_out"]
    return [
        ("annotations counted", int(r.stats.n)),
        ("median (%)", _pct(r.stats.median)),
        ("mean (%)", _pct(r.stats.mean)),
        ("SEM (%)", _pct(r.stats.sem)),
        ("mode, bin center (%)", _pct(r.mode)),
        ("outliers omitted", int(r.n_omitted_outlier)),
        ("outside the surfaces", int(r.n_omitted_nan)),
        ("soma left out", int((left == "soma").sum())),
    ]


# -- the workbook ------------------------------------------------------------------------------------

def _sheet_names(cells: Sequence[StrataCell]) -> list[str]:
    """A sheet name per cell, at most 31 characters and unique (Excel's rules)."""
    names: list[str] = []
    for cell in cells:
        base = f"{cell.name} {cell.volume}"
        for bad in '[]:*?/\\':
            base = base.replace(bad, " ")
        name = base[:31]
        n = 2
        while name.lower() in {x.lower() for x in names} or name.lower() == "summary":
            suffix = f" ({n})"
            name = base[:31 - len(suffix)] + suffix
            n += 1
        names.append(name)
    return names


def stratification_xlsx(cells: Sequence[StrataCell], settings: StrataSettings) -> bytes:
    """The chart's data as an Excel workbook: a Summary sheet with a column per
    cell (each bin's center and count, the statistics, the settings) and a sheet
    per cell with every annotation's depth."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    bold = Font(bold=True)
    muted = Font(italic=True, color="666666")

    summary.append(["IPL stratification"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([settings_line(settings)])
    summary["A2"].font = muted
    summary.append([])

    # Two columns a cell: its bins' centers and counts, matched by bin number.
    head = ["bin"]
    for cell in cells:
        head += [f"{cell.label} center (%)", f"{cell.label} count"]
    summary.append(head)
    for col in range(1, len(head) + 1):
        summary.cell(row=summary.max_row, column=col).font = bold
    bins = max((len(c.result.centers) for c in cells), default=0)
    for i in range(bins):
        row: list = [i + 1]
        for cell in cells:
            if i < len(cell.result.centers):
                row += [_pct(cell.result.centers[i]), int(cell.result.counts[i])]
            else:
                row += [None, None]
        summary.append(row)
    summary.append([])

    # The statistics, under the same columns: a cell's value under its name, in its first column.
    stats_head = summary.max_row + 1
    summary.append(["statistic"] + sum(([cell.label, None] for cell in cells), []))
    for col in range(1, len(head) + 1):
        summary.cell(row=stats_head, column=col).font = bold
    rows = [_stats_rows(cell) for cell in cells]
    for k, (name, _) in enumerate(rows[0] if rows else []):
        summary.append([name] + sum(([r[k][1], None] for r in rows), []))
    summary.append(["volume"] + sum(([cell.volume, None] for cell in cells), []))
    summary.append([])

    summary.append(["settings"])
    summary.cell(row=summary.max_row, column=1).font = bold
    for name, value in [
        ("bins per cell", settings.num_bins),
        ("soma", "included" if settings.include_soma else "left out (over 80 % of the largest radius)"),
        ("outliers", f"omitted beyond {settings.ipl_margin * 100:g} % of the IPL" if settings.omit_outliers
         else "kept"),
        ("gridlines every (%)", settings.tick_pct),
        ("depth", "percent of the IPL: 0 at the INL border, 100 at the GCL border"),
        ("depth measured", METHOD_WORDS.get(settings.method, settings.method)
         + ("; depth (µm) from the INL border and the IPL's thickness through each annotation "
            "the same way")),
        ("bins", "each cell's own: its range cut into equal bins (numpy.histogram); matched here by number"),
        ("SBFSEM-tools", settings.version),
    ]:
        summary.append([name, value])
    summary.column_dimensions["A"].width = 22
    for col in range(2, len(head) + 1):
        summary.column_dimensions[get_column_letter(col)].width = 16

    columns = {
        "location_id": "Location ID", "section": "section", "x_um": "X (µm)", "y_um": "Y (µm)",
        "z_um": "Z (µm)", "radius_um": "radius (µm)", "ipl_depth": "IPL depth (%)",
        "depth_um": "depth (µm)", "ipl_thickness_um": "IPL thickness (µm)",
        "counted": "counted", "left_out": "left out because",
    }
    for cell, name in zip(cells, _sheet_names(cells), strict=True):
        sheet = book.create_sheet(name)
        table = cell.annotations.copy()
        table["ipl_depth"] = [_pct(v) for v in table["ipl_depth"]]
        table = table[list(columns)].rename(columns=columns)
        sheet.append([f"{cell.label} · {cell.volume} · {METHOD_WORDS.get(settings.method, settings.method)}"])
        sheet["A1"].font = bold
        sheet.append(list(table.columns))
        for col in range(1, len(table.columns) + 1):
            sheet.cell(row=2, column=col).font = bold
        for values in table.itertuples(index=False):
            sheet.append([_cell_value(v) for v in values])
        sheet.freeze_panes = "A3"
        for col, width in enumerate([13, 9, 11, 11, 11, 11, 13, 11, 17, 9, 18], start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _cell_value(value):
    """A value as Excel takes it: plain Python numbers, None for NaN."""
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return float(value) if math.isfinite(float(value)) else None
    if isinstance(value, np.bool_):
        return bool(value)
    return value


# -- the PDF -------------------------------------------------------------------------------------------

#: The chart's page, inches: the figure and a line under it.
CHART_PAGE = (8.0, 4.6)


def _chart_page(cells: Sequence[StrataCell], settings: StrataSettings):
    from matplotlib.figure import Figure

    from .pdfpage import INK, MUTED

    fig = Figure(figsize=CHART_PAGE)
    fig.patch.set_facecolor("white")
    ax = fig.add_axes((0.085, 0.2, 0.89, 0.74))
    usable = [c for c in cells if len(c.result.centers)]
    centers = np.concatenate([c.result.centers for c in usable]) if usable else np.empty(0)
    counts = np.concatenate([c.result.counts for c in usable]) if usable else np.empty(0)
    x0 = min(-MARGIN_SHOWN, float(centers.min())) if centers.size else -MARGIN_SHOWN
    x1 = max(1 + MARGIN_SHOWN, float(centers.max())) if centers.size else 1 + MARGIN_SHOWN
    y1 = max(1.0, float(counts.max())) if counts.size else 1.0

    # Beyond the IPL, shaded; the gridlines and labels from the INL to the GCL.
    for a, b in ((x0, 0.0), (1.0, x1)):
        ax.axvspan(a * 100, b * 100, color="#000000", alpha=0.05, lw=0)
    ticks = ipl_ticks(settings.tick_pct)
    for value, _ in ticks:
        ax.axvline(value * 100, color="#000000", alpha=0.15, lw=0.6, ls=(0, (3, 3)), zorder=0)
    # Labels no closer than 20 pt (the page's axes are 0.89 of 8 in across).
    points_per_fraction = 0.89 * CHART_PAGE[0] * 72 / (x1 - x0)
    shown = labeled_ticks(ticks, (ticks[1][0] - ticks[0][0]) * points_per_fraction, 20.0)
    ax.set_xticks([v * 100 for v, _ in ticks], [label if i in shown else "" for i, (_, label) in enumerate(ticks)])
    ax.set_xlim(x0 * 100, x1 * 100)
    ax.set_ylim(0, y1 * 1.04)
    ax.tick_params(labelsize=8, colors=MUTED, length=3, width=0.6)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_linewidth(0.6)
        ax.spines[side].set_color("#9aa1ab")
    ax.set_xlabel("IPL depth (%)", fontsize=9, color=INK)
    ax.set_ylabel("Annotations", fontsize=9, color=INK)

    for cell in usable:
        x = np.asarray(cell.result.centers, dtype=float) * 100
        ax.plot(x, cell.result.counts, color=cell.color, lw=1.4, marker="o", ms=3.2,
                solid_joinstyle="round", label=cell.label)
        median = cell.result.stats.median
        if median is not None and math.isfinite(float(median)):
            ax.plot([float(median) * 100], [y1 * 0.06], marker="v", ms=7, color=cell.color,
                    alpha=0.85, ls="none")
    if usable:
        legend = ax.legend(loc="upper right", fontsize=8, frameon=False, handlelength=1.2)
        for text in legend.get_texts():
            text.set_color(INK)
    else:
        ax.text(50, y1 / 2, "No annotations fall within the fitted IPL boundary surfaces.",
                ha="center", va="center", fontsize=9, color=MUTED)
    fig.text(0.085, 0.045, settings_line(settings), fontsize=6.5, color=MUTED, ha="left")
    return fig


def annotation_rows_per_page(page_height_in: float = 11.0) -> int:
    """How many annotation rows a data page holds (the page estimate the dialog shows)."""
    row_h = 7.2 * 1.62 / 72.0
    top = 0.42 + 14 * 1.35 / 72.0 + 8.5 * 1.5 / 72.0 + 0.2 + row_h
    return int((page_height_in - 0.6 - top) / row_h)


def _data_pages(cells: Sequence[StrataCell], settings: StrataSettings, *, statistics: bool,
                annotations: bool, first_number: int, footer: str) -> list:
    from .pdfpage import PAGE_SIZES, Column, Page

    size = PAGE_SIZES["letter"]
    pages: list[Page] = []

    def new_page(title: str, subtitle: str = "") -> tuple[Page, float]:
        page = Page(size, first_number + len(pages), footer=footer)
        pages.append(page)
        return page, page.header(title, subtitle)

    if statistics:
        page, y = new_page("IPL stratification: statistics", settings_line(settings))
        names = [name for name, _ in _stats_rows(cells[0])] if cells else []
        columns = [Column("cell", 1.9), Column("volume", 1.25)] + [
            Column(n.replace(" (%)", " %").replace(", bin center", ""), 0.53, "right") for n in names]
        scale = min(1.0, page.usable / sum(c.width for c in columns))
        columns = [Column(c.title, c.width * scale, c.align) for c in columns]
        rows: list[list[Any]] = [[c.label, c.volume] + ["" if v is None else v for _, v in _stats_rows(c)]
                                 for c in cells]
        y = page.table(page.left, y, columns, rows, size=6.6, swatches=[c.color for c in cells])
        y += 0.25
        # The binned counts, five cells a block: the bin, then each cell's center and count.
        for start in range(0, len(cells), 5):
            group = cells[start:start + 5]
            bins = max((len(c.result.centers) for c in group), default=0)
            need = (bins + 2) * 7.2 * 1.62 / 72.0 + 0.3
            if y + need > page.bottom:
                page, y = new_page("IPL stratification: binned counts")
            y = page.section(y, "Binned counts" + (f" (cells {start + 1} to {start + len(group)})"
                                                   if len(cells) > 5 else ""))
            cols = [Column("bin", 0.4, "right")]
            for c in group:
                cols += [Column(f"{c.name} center %", 0.85, "right"), Column("count", 0.55, "right")]
            rows = []
            for i in range(bins):
                row: list = [i + 1]
                for c in group:
                    if i < len(c.result.centers):
                        row += [f"{_pct(c.result.centers[i]):.2f}", int(c.result.counts[i])]
                    else:
                        row += ["", ""]
                rows.append(row)
            y = page.table(page.left, y, cols, rows, size=6.6) + 0.25

    if annotations:
        cols = [Column("Location ID", 0.85, "right"), Column("section", 0.55, "right"),
                Column("X µm", 0.7, "right"), Column("Y µm", 0.7, "right"), Column("Z µm", 0.7, "right"),
                Column("radius µm", 0.7, "right"), Column("IPL depth %", 0.75, "right"),
                Column("depth µm", 0.75, "right"), Column("IPL µm", 0.75, "right"),
                Column("left out because", 0.95)]

        def um(value) -> str:
            return f"{value:.3f}" if math.isfinite(float(value)) else ""

        for cell in cells:
            table = cell.annotations
            rows = [[int(r.location_id), int(r.section) if math.isfinite(float(r.section)) else "",
                     f"{r.x_um:.3f}", f"{r.y_um:.3f}", f"{r.z_um:.3f}", f"{r.radius_um:.3f}",
                     "" if _pct(r.ipl_depth) is None else f"{_pct(r.ipl_depth):.2f}",
                     um(r.depth_um), um(r.ipl_thickness_um), r.left_out]
                    for r in table.itertuples(index=False)]
            dim = [None if r[-1] == "" else "#8a919b" for r in rows]
            done = 0
            part = 1
            while done < len(rows) or done == 0:
                page, y = new_page(f"IPL depth of each annotation: {cell.label}",
                                   f"{cell.volume} · {len(rows):,} annotations, "
                                   f"{int(table['counted'].sum()):,} counted · "
                                   f"{METHOD_WORDS.get(settings.method, settings.method)}"
                                   + (f" · part {part}" if part > 1 else ""))
                per = annotation_rows_per_page(size[1])
                chunk = rows[done:done + per]
                page.table(page.left, y, cols, chunk, size=7.2, colors=dim[done:done + per])
                done += max(len(chunk), 1)
                part += 1
                if not rows:
                    break
    for page in pages:
        page.finish()
    return pages


def stratification_pdf(cells: Sequence[StrataCell], settings: StrataSettings, *,
                       statistics: bool = False, annotations: bool = False) -> bytes:
    """The chart as a vector PDF (its first page, the size of the figure), and
    with ``statistics`` and ``annotations`` the data on the pages after it."""
    from matplotlib.backends.backend_pdf import PdfPages

    footer = "IPL stratification · " + ", ".join(f"{c.name} ({c.volume})" for c in cells)
    out = io.BytesIO()
    with PdfPages(out, metadata={
        "Title": "IPL stratification",
        "Subject": ", ".join(f"{c.name} ({c.volume})" for c in cells),
        "Creator": f"SBFSEM-tools {settings.version}".strip(),
    }) as pdf:
        pdf.savefig(_chart_page(cells, settings))
        if statistics or annotations:
            for page in _data_pages(cells, settings, statistics=statistics, annotations=annotations,
                                    first_number=2, footer=footer):
                pdf.savefig(page.fig)
    return out.getvalue()
